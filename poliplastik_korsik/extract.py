"""Extraction and source locations for the POLIPLASTIK PDF/PPTX corpus."""
from __future__ import annotations

import io
import logging
import math
from pathlib import Path

import pymupdf
from lxml import etree
from PIL import Image
from pptx import Presentation

from ocr_core import image_rgb, ocr_reasons, sha, tidy, write_json

LOG = logging.getLogger(__name__)
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'dgm': 'http://schemas.openxmlformats.org/drawingml/2006/diagram',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}


def smartart(blob: bytes) -> list[dict]:
    root = etree.fromstring(blob, parser=etree.XMLParser(resolve_entities=False, no_network=True))
    nodes = {}
    for p in root.findall('./dgm:ptLst/dgm:pt', NS):
        if p.get('type', 'node') not in ('node', 'doc'):
            continue
        text = tidy('\n'.join(''.join(t.itertext()) for t in p.findall('./dgm:t/a:p', NS)))
        nodes[p.get('modelId')] = {'text': text, 'type': p.get('type', 'node')}
    children, incoming = {}, set()
    for c in root.findall('./dgm:cxnLst/dgm:cxn', NS):
        if c.get('type', 'parOf') != 'parOf':
            continue
        src, dst = c.get('srcId'), c.get('destId')
        if src in nodes and dst in nodes:
            children.setdefault(src, []).append((int(c.get('srcOrd', '0')), dst))
            incoming.add(dst)
    result, visited = [], set()
    def walk(key, path):
        if key in visited:
            return
        visited.add(key)
        node = nodes[key]
        current = path + [node['text']] if node['text'] else path
        if node['text']:
            result.append({'text': node['text'], 'node_id': key, 'ancestors': path})
        for _, child in sorted(children.get(key, [])):
            walk(child, current)
    for key in nodes:
        if key not in incoming:
            walk(key, [])
    for key in nodes:
        walk(key, [])
    return result


def segment(source: dict, kind: str, number: int) -> dict:
    return {'document': source['name'], 'source_path': source['path'],
            'source_sha256': source['sha256'], 'location_type': kind, 'location': number,
            'segment_id': f"{source['sha256']}:{kind}:{number}",
            'blocks': [], 'artifacts': [], 'review_reasons': [], 'warnings': [],
            'ocr_attempted': False, 'ocr_mean_confidence': None,
            'extraction_method': '', 'service_content': False}


def finish(seg: dict):
    seg['text'] = tidy('\n\n'.join(b['text'] for b in seg['blocks'] if b.get('text')))
    seg['raw_text'] = seg['text']
    seg['char_count'] = len(seg['text'])
    seg['review_reasons'] = sorted(set(seg['review_reasons']))
    seg['needs_review'] = bool(seg['review_reasons'])
    seg['status'] = 'excluded' if seg['service_content'] else 'needs_review' if seg['needs_review'] else 'accepted'
    seg['content_fingerprint'] = sha(__import__('json').dumps(
        [seg['text'], seg['blocks'], seg['review_reasons'], seg['service_content']], ensure_ascii=False, sort_keys=True))
    return seg


def apply_light_policy(seg: dict):
    """Advisory OCR warnings, hard failures still block; do not label human-reviewed."""
    if seg['service_content']:
        return
    seg['quality_warnings'] = list(seg['review_reasons'])
    hard = {'empty_page', 'empty_slide', 'empty_smartart', 'empty_ocr',
            'very_little_page_text', 'image_ocr_error', 'unsupported_embedded_object',
            'unsupported_graphic_frame', 'unreadable_image_or_illustration'}
    failed = hard.intersection(seg['review_reasons'])
    for artifact in seg['artifacts']:
        if artifact['kind'] == 'page' and (artifact['confidence'] or 0) < 50:
            failed.add('unreadable_main_page')
    if not seg['text'].strip():
        failed.add('no_main_text')
    seg['blocking_reasons'] = sorted(failed)
    seg['needs_review'] = bool(failed)
    seg['status'] = 'needs_review' if failed else 'accepted'


def add_ocr(seg: dict, image: Image.Image, ocr, assets: Path | None, key: str, kind='image', bbox=None):
    image = image_rgb(image)
    if assets:
        image_path = assets/(key + '.png')
        image.save(image_path)
    result = ocr.recognize(image, kind)
    if assets:
        evidence_path = assets/(key + '.ocr.json')
        write_json(evidence_path, result)
    artifact = {'kind': kind, 'image': 'assets/' + image_path.name if assets else None,
                'ocr_evidence': 'assets/' + evidence_path.name if assets else None,
                'text': result['text'], 'confidence': result['confidence'],
                'bbox': bbox, 'original_size': list(image.size),
                'review_reasons': ocr_reasons(result, kind)}
    seg['artifacts'].append(artifact)
    seg['ocr_attempted'] = True
    seg['review_reasons'].extend(artifact['review_reasons'])
    if result['text']:
        # Tiny logo/photo OCR is kept as evidence but not as a knowledge paragraph.
        useful = kind == 'page' or sum(c.isalpha() for c in result['text']) >= 12
        if useful:
            seg['blocks'].append({'kind': 'ocr', 'text': result['text'], 'bbox': bbox,
                                  'artifact': artifact['image'], 'confidence': result['confidence']})
    scores = [a['confidence'] for a in seg['artifacts'] if a['text'] and a['confidence'] is not None]
    seg['ocr_mean_confidence'] = round(sum(scores)/len(scores), 2) if scores else None
    return result


def pdf_segments(source: dict, raw: bytes, args, ocr, assets: Path):
    result = []
    with pymupdf.open(stream=raw, filetype='pdf') as pdf:
        if pdf.needs_pass:
            raise ValueError('PDF зашифрован: требуется версия без пароля')
        if len(pdf) > args.max_pages:
            raise ValueError('Превышен --max-pages')
        for number, page in enumerate(pdf, 1):
            LOG.info('%s: страница %d/%d', source['name'], number, len(pdf))
            seg = segment(source, 'page', number)
            non_knowledge = getattr(args, 'non_knowledge_segments', {}).get(seg['segment_id'])
            if non_knowledge and not getattr(args, 'audit', False):
                seg['service_content'] = True
                seg['extraction_method'] = 'excluded_non_knowledge'
                seg['warnings'].append(non_knowledge)
                result.append(finish(seg))
                continue
            native = tidy(page.get_text('text', sort=True))
            good_native = len(native) >= 80 and '\ufffd' not in native and sum(c.isalpha() for c in native)/max(1,len(native)) > .45
            if good_native:
                for block in page.get_text('blocks', sort=True):
                    if block[6] == 0 and tidy(block[4]):
                        seg['blocks'].append({'kind': 'pdf_text', 'text': tidy(block[4]), 'bbox': list(block[:4])})
                seg['extraction_method'] = 'pdf_text'
                # Native text alone is insufficient for a mixed text-and-scan PDF.
                for idx, info in enumerate(page.get_image_info()):
                    rect = pymupdf.Rect(info['bbox']) & page.rect
                    if rect.is_empty or rect.get_area()/page.rect.get_area() < .025:
                        continue
                    # Existing text in the image rectangle normally means a searchable scan.
                    inside = page.get_text('text', clip=rect)
                    if len(inside.strip()) > 80:
                        continue
                    if rect.width*rect.height*(args.dpi/72)**2 > args.max_pixels:
                        raise ValueError('PDF image too large for --max-pixels')
                    pix = page.get_pixmap(dpi=args.dpi, clip=rect, alpha=False)
                    image = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
                    add_ocr(seg, image, ocr, assets, f"{source['sha256'][:16]}_p{number:03}_i{idx}", bbox=list(rect))
                if seg['ocr_attempted']:
                    seg['extraction_method'] += '+image_ocr'
                    seg['review_reasons'].append('mixed_pdf_layout')
            else:
                pixels = page.rect.width*page.rect.height*(args.dpi/72)**2
                if not math.isfinite(pixels) or pixels > args.max_pixels:
                    raise ValueError('PDF page too large for --max-pixels')
                pix = page.get_pixmap(dpi=args.dpi, alpha=False)
                image = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
                seg['native_text_before_ocr'] = native
                add_ocr(seg, image, ocr, assets, f"{source['sha256'][:16]}_p{number:03}", kind='page', bbox=list(page.rect))
                seg['extraction_method'] = 'ocr'
            if not seg['blocks']:
                seg['review_reasons'].append('empty_page')
            result.append(finish(seg))
    return result


def pptx_segments(source: dict, raw: bytes, args, ocr, assets: Path):
    prs = Presentation(io.BytesIO(raw))
    if len(prs.slides) > args.max_pages:
        raise ValueError('Превышен --max-pages')
    result = []
    for number, slide in enumerate(prs.slides, 1):
        LOG.info('%s: слайд %d/%d', source['name'], number, len(prs.slides))
        seg = segment(source, 'slide', number)
        seg['slide_size_emu'] = [prs.slide_width, prs.slide_height]
        seg['hidden_slide'] = slide._element.get('show') == '0'
        seg['extraction_method'] = 'pptx_text'
        native_chars = sum(len(t) for t in slide._element.xpath('.//a:t/text()'))
        quiz = any('ISPRING_QUIZ' in el.get('name', '').upper() for el in slide._element.xpath('.//p:cNvPr'))
        if quiz and not getattr(args, 'audit', False):
            seg['service_content'] = True
            seg['extraction_method'] = 'excluded_quiz_placeholder'
            seg['warnings'].append('iSpring quiz placeholder; actual quiz questions are not available in slide text')
            result.append(finish(seg))
            continue
        seen_diagrams = set()
        def walk(shapes):
            ordered = sorted(shapes, key=lambda x: (x.top, x.left))
            for shape in ordered:
                bbox = [shape.left, shape.top, shape.width, shape.height]
                if shape.shape_type == 6:  # Group, with its own local coordinate system.
                    walk(shape.shapes)
                    continue
                if shape.has_text_frame and tidy(shape.text):
                    seg['blocks'].append({'kind': 'pptx_text', 'text': tidy(shape.text),
                                          'shape_id': shape.shape_id, 'bbox_emu': bbox})
                if shape.has_table:
                    rows = [' | '.join(tidy(c.text) for c in row.cells if not c.is_spanned) for row in shape.table.rows]
                    seg['blocks'].append({'kind': 'table', 'text': '\n'.join(rows), 'rows': rows,
                                          'shape_id': shape.shape_id, 'bbox_emu': bbox})
                # SmartArt is a graphicFrame, not a picture or text frame.
                relids = shape._element.findall('.//dgm:relIds', NS)
                for relid in relids:
                    rid = relid.get('{%s}dm' % NS['r'])
                    if not rid or rid in seen_diagrams:
                        continue
                    seen_diagrams.add(rid)
                    nodes = smartart(slide.part.related_part(rid).blob)
                    text = '\n\n'.join((' / '.join(n['ancestors'])+'\n' if n['ancestors'] else '') + n['text'] for n in nodes)
                    seg['blocks'].append({'kind': 'smartart', 'text': text, 'nodes': nodes,
                                          'shape_id': shape.shape_id, 'bbox_emu': bbox})
                    if not nodes:
                        seg['review_reasons'].append('empty_smartart')
                if shape.shape_type == 13:
                    try:
                        classification = getattr(args, 'image_policy', {}).get(sha(shape.image.blob))
                        if not getattr(args, 'audit', False) and classification and (
                            classification.get('role') == 'illustration' or
                            classification.get('role') == 'document_image' and native_chars >= 80
                        ):
                            seg['artifacts'].append({'kind':'image','image':None,'ocr_evidence':None,
                                'text':'','confidence':None,'bbox':bbox,'classification':classification,
                                'source_image_sha256':sha(shape.image.blob),'review_reasons':[],
                                'skipped_reason':'illustration' if classification['role']=='illustration' else 'supporting_screenshot_native_text_retained'})
                            continue
                        image = Image.open(io.BytesIO(shape.image.blob))
                        if image.width*image.height > getattr(args, 'max_pixels', 40_000_000):
                            raise ValueError('Embedded image exceeds --max-pixels')
                        image = image_rgb(image)
                        w, h = image.size
                        # PowerPoint stores full images, including cropped-out text.
                        box = (round(w*shape.crop_left), round(h*shape.crop_top),
                               round(w*(1-shape.crop_right)), round(h*(1-shape.crop_bottom)))
                        if box[2] <= box[0] or box[3] <= box[1]:
                            raise ValueError('invalid picture crop')
                        image = image.crop(box)
                        key = f"{source['sha256'][:16]}_s{number:03}_sh{shape.shape_id}"
                        classification = getattr(args, 'image_policy', {}).get(sha(shape.image.blob))
                        if classification and classification.get('role') == 'illustration':
                            image.save(assets/(key+'.png'))
                            seg['artifacts'].append({'kind':'image','image':'assets/'+key+'.png',
                                'ocr_evidence':None,'text':'','confidence':None,'bbox':bbox,
                                'classification':classification,'source_image_sha256':sha(shape.image.blob),
                                'review_reasons':[]})
                        else:
                            extracted = add_ocr(seg, image, ocr, assets,key,bbox=bbox)
                            seg['artifacts'][-1]['source_image_sha256'] = sha(shape.image.blob)
                            if classification and classification.get('role') == 'document_image':
                                seg['artifacts'][-1]['classification'] = classification
                                seg['review_reasons'].append('document_image_requires_verification')
                            if len(extracted['text']) < 12 and max(image.size) > 500:
                                seg['review_reasons'].append('unreadable_image_or_illustration')
                    except Exception as exc:
                        seg['review_reasons'].append('image_ocr_error')
                        seg['warnings'].append(f'shape {shape.shape_id}: {type(exc).__name__}: {exc}')
                elif shape.shape_type in (7, 10, 11, 12, 16):
                    seg['review_reasons'].append('unsupported_embedded_object')
                elif shape.shape_type is None and not relids and not shape.has_table:
                    seg['review_reasons'].append('unsupported_graphic_frame')
                if 'ISPRING_QUIZ' in shape.name.upper():
                    seg['service_content'] = True
                    seg['warnings'].append('iSpring quiz placeholder; actual quiz questions are not available in slide text')
        walk(slide.shapes)
        seg['links'] = [r.target_ref for r in slide.part.rels.values() if r.is_external]
        if seg['hidden_slide']:
            seg['service_content'] = True
            seg['warnings'].append('hidden_slide')
        if not seg['blocks']:
            seg['review_reasons'].append('empty_slide')
        if seen_diagrams:
            seg['extraction_method'] += '+smartart'
        if seg['ocr_attempted']:
            seg['extraction_method'] += '+image_ocr'
        result.append(finish(seg))
    return result
