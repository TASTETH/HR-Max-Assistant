import argparse
import io
import json
import random
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from chunks import split_chunks
from extract import apply_light_policy, finish, segment, smartart, pptx_segments
from ingest import checked_zip, load_current, output_lock, parser, run
from ocr_core import parse_tsv, sha, tidy, write_json
from review import apply_reviews, write_review
from pptx import Presentation
from pptx.util import Inches
from PIL import Image


SOURCE={'name':'test.pptx','path':'folder/test.pptx','sha256':'a'*64}


def sample(text='Достоверный текст документа.'):
    s=segment(SOURCE,'slide',1)
    s['blocks']=[{'kind':'pptx_text','text':text}]
    return finish(s)


class IngestionTests(unittest.TestCase):
    def test_smartart_hierarchy_does_not_mix_categories(self):
        # Storage order differs from display/model parent order.
        blob=b'''<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><dgm:ptLst><dgm:pt modelId="root" type="doc"/><dgm:pt modelId="a"><dgm:t><a:p><a:r><a:t>Category A</a:t></a:r></a:p></dgm:t></dgm:pt><dgm:pt modelId="b"><dgm:t><a:p><a:r><a:t>Category B</a:t></a:r></a:p></dgm:t></dgm:pt><dgm:pt modelId="child"><dgm:t><a:p><a:r><a:t>Rule A1</a:t></a:r></a:p></dgm:t></dgm:pt></dgm:ptLst><dgm:cxnLst><dgm:cxn srcId="root" destId="a" srcOrd="0"/><dgm:cxn srcId="root" destId="b" srcOrd="1"/><dgm:cxn srcId="a" destId="child" srcOrd="0"/><dgm:cxn type="presOf" srcId="b" destId="child"/></dgm:cxnLst></dgm:dataModel>'''
        nodes=smartart(blob)
        self.assertEqual([n['text'] for n in nodes],['Category A','Rule A1','Category B'])
        self.assertEqual(nodes[1]['ancestors'],['Category A'])

    def test_chunk_offsets_bound_and_coverage(self):
        rng=random.Random(12)
        for _ in range(40):
            text='\n\n'.join(' '.join('слово'+str(rng.randint(0,100)) for _ in range(rng.randint(3,90)))+'.' for _ in range(20))
            seg=sample(text);chunks=split_chunks(seg,400,70)
            covered=set()
            for c in chunks:
                self.assertLessEqual(len(c['text']),400)
                self.assertEqual(c['text'],seg['text'][c['char_start']:c['char_end']])
                covered.update(range(c['char_start'],c['char_end']))
            self.assertTrue(all(i in covered or ch.isspace() for i,ch in enumerate(seg['text'])))
            self.assertEqual(len({c['chunk_id'] for c in chunks}),len(chunks))

    def test_pending_and_excluded_never_become_chunks(self):
        s=sample()
        for state in ['needs_review','excluded']:
            s['status']=state
            self.assertEqual(split_chunks(s),[])

    def test_chunk_content_changes_id(self):
        a=split_chunks(sample('Первый вариант текста.'))[0]
        b=split_chunks(sample('Второй вариант текста.'))[0]
        self.assertNotEqual(a['chunk_id'],b['chunk_id'])

    def test_reviews_are_hash_bound_and_audited(self):
        with tempfile.TemporaryDirectory() as tmp:
            s=sample();s['status']='needs_review';s['needs_review']=True;s['review_reasons']=['low_ocr_confidence']
            d={k:s[k] for k in ['segment_id','source_sha256','content_fingerprint']}
            d.update(decision='approve',text='Текст после сверки.')
            path=Path(tmp)/'reviews.json';write_json(path,{'schema_version':1,'decisions':[d]})
            apply_reviews([s],path)
            self.assertEqual(s['status'],'reviewed');self.assertEqual(s['raw_text'],'Достоверный текст документа.')
            self.assertTrue(s['review']['text_changed'])
            d['content_fingerprint']='stale';write_json(path,{'schema_version':1,'decisions':[d]})
            with self.assertRaises(ValueError):apply_reviews([sample()],path)

    def test_review_cannot_approve_service_stub(self):
        with tempfile.TemporaryDirectory() as tmp:
            s=sample();s['service_content']=True
            d={k:s[k] for k in ['segment_id','source_sha256','content_fingerprint']};d['decision']='approve'
            path=Path(tmp)/'reviews.json';write_json(path,{'schema_version':1,'decisions':[d]})
            with self.assertRaises(ValueError):apply_reviews([s],path)

    def test_zip_path_and_size_limits(self):
        args=argparse.Namespace(max_total_mb=1,max_file_mb=1)
        for name in ['../escape.pdf','/absolute.pdf','C:/bad.pdf']:
            b=io.BytesIO()
            with zipfile.ZipFile(b,'w') as z:z.writestr(name,b'x')
            with zipfile.ZipFile(io.BytesIO(b.getvalue())) as z:
                with self.assertRaises(ValueError):checked_zip(z,args)

    def test_output_lock_is_exclusive(self):
        with tempfile.TemporaryDirectory() as tmp:
            with output_lock(Path(tmp)):
                with self.assertRaises(RuntimeError):
                    with output_lock(Path(tmp)):pass
            # A persistent file is safe: the OS releases ownership on close/crash.
            with output_lock(Path(tmp)):pass

    def test_snapshot_gate_and_integrity(self):
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp);snap=out/'runs'/'test';snap.mkdir(parents=True)
            report=snap/'report.json';write_json(report,{'ready_for_embeddings':False})
            write_json(out/'current.json',{'snapshot':'runs/test','files':{'report.json':sha(report.read_bytes())}})
            with self.assertRaises(RuntimeError):load_current(out)
            self.assertEqual(load_current(out,False)[0],snap)
            report.write_text('{}')
            with self.assertRaises(ValueError):load_current(out,False)

    def test_tsv_preserves_paragraphs_and_word_evidence(self):
        t='level\tblock_num\tpar_num\tline_num\tleft\ttop\twidth\theight\tconf\ttext\n5\t1\t1\t1\t10\t20\t30\t12\t90\tТекст\n5\t1\t2\t1\t10\t50\t20\t12\t80\t4384\n'
        r=parse_tsv(t,2)
        self.assertEqual(r['text'],'Текст\n\n4384')
        self.assertEqual(r['words'][0]['bbox'],[5,10,15,6])

    def test_non_tsv_output_cannot_silently_pass(self):
        with self.assertRaises(ValueError):
            parse_tsv('Just plain text instead of TSV')

    def test_tesseract_options_precede_output_config(self):
        from ocr_core import OCR
        ocr=OCR.__new__(OCR);ocr.executable='tesseract';ocr.tessdata='my data';ocr.timeout=5
        completed=argparse.Namespace(stdout='',stderr='')
        with patch('ocr_core.subprocess.run',return_value=completed) as runner:
            ocr.command(['input.png','stdout','-c','tessedit_create_tsv=1'])
        self.assertEqual(runner.call_args.args[0][:3],['tesseract','--tessdata-dir','my data'])

    def test_partial_language_failure_is_fatal_even_exit_zero(self):
        from ocr_core import OCR
        ocr=OCR.__new__(OCR);ocr.executable='tesseract';ocr.tessdata=None;ocr.timeout=5
        completed=argparse.Namespace(stdout='some English text',stderr="Failed loading language 'rus'")
        with patch('ocr_core.subprocess.run',return_value=completed):
            with self.assertRaises(RuntimeError):ocr.command(['input.png','stdout'])

    def test_native_group_and_table_extraction(self):
        prs=Presentation();s=prs.slides.add_slide(prs.slide_layouts[6])
        group=s.shapes.add_group_shape();box=group.shapes.add_textbox(0,0,Inches(3),Inches(1));box.text='Текст внутри группы'
        table=s.shapes.add_table(2,2,0,Inches(2),Inches(5),Inches(2)).table
        table.cell(0,0).text='Параметр';table.cell(0,1).text='Значение';table.cell(1,0).text='Срок';table.cell(1,1).text='3 дня'
        buf=io.BytesIO();prs.save(buf)
        with tempfile.TemporaryDirectory() as tmp:
            segs=pptx_segments(SOURCE,buf.getvalue(),argparse.Namespace(max_pages=10),None,Path(tmp))
        self.assertIn('Текст внутри группы',segs[0]['text'])
        self.assertIn('Срок | 3 дня',segs[0]['text'])

    def test_review_html_escapes_document_script(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'review.html';write_review(path,[sample('</script><script>alert(1)</script>')],{})
            text=path.read_text()
            self.assertNotIn('</script><script>alert',text)
            self.assertIn('\\u003c/script>',text)

    def test_failed_document_commits_blocked_report(self):
        with tempfile.TemporaryDirectory() as tmp:
            inp=Path(tmp)/'broken.pdf';inp.write_bytes(b'not a PDF')
            args=parser().parse_args(['--input',str(inp),'--output',str(Path(tmp)/'out')])
            fake=argparse.Namespace(signature={'test':True},calls=0,hits=0)
            with patch('ingest.OCR',return_value=fake), self.assertLogs('ingest',level='ERROR'):
                report=run(args)
            self.assertFalse(report['ready_for_embeddings']);self.assertEqual(report['failed_documents'],1)
            with self.assertRaises(RuntimeError):load_current(args.output)

    def test_tidy_does_not_change_legal_numbers(self):
        self.assertEqual(tidy('4384\v71-4384\u00a0 3000,00\u00ad'),'4384\n71-4384 3000,00')

    def test_main_text_warnings_are_advisory_but_empty_content_blocks(self):
        s=sample();s['review_reasons']=['low_ocr_confidence'];apply_light_policy(s)
        self.assertEqual(s['status'],'accepted');self.assertFalse(s['needs_review'])
        self.assertEqual(s['quality_warnings'],['low_ocr_confidence'])
        s=sample('');s['review_reasons']=['empty_page'];apply_light_policy(s)
        self.assertEqual(s['status'],'needs_review');self.assertIn('empty_page',s['blocking_reasons'])

    def test_native_slide_kept_without_ocr_of_known_supporting_image(self):
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        box=slide.shapes.add_textbox(0,0,Inches(6),Inches(2))
        box.text='Основная инструкция: отправить уведомление руководителю. '*4
        data=io.BytesIO();Image.new('RGB',(100,100),'white').save(data,format='PNG');raw=data.getvalue()
        slide.shapes.add_picture(io.BytesIO(raw),0,Inches(3),Inches(2),Inches(2))
        deck=io.BytesIO();prs.save(deck)
        args=argparse.Namespace(max_pages=10,audit=False,image_policy={sha(raw):{'role':'document_image'}})
        seg=pptx_segments(SOURCE,deck.getvalue(),args,None,None)[0]
        self.assertIn('Основная инструкция',seg['text']);self.assertFalse(seg['ocr_attempted'])
        self.assertEqual(seg['artifacts'][0]['skipped_reason'],'supporting_screenshot_native_text_retained')

    def test_main_mode_does_not_write_png_or_ocr_evidence(self):
        from extract import add_ocr
        fake=argparse.Namespace(recognize=lambda im,kind:{'text':'Хорошо распознанный текст документа.',
            'confidence':95,'low_confidence_ratio':0,'median_word_height':15,'attempts':[]})
        s=sample();add_ocr(s,Image.new('RGB',(30,30)),fake,None,'unused',kind='page')
        self.assertIsNone(s['artifacts'][0]['image']);self.assertIsNone(s['artifacts'][0]['ocr_evidence'])
        self.assertIn('Хорошо распознанный',s['blocks'][-1]['text'])


if __name__=='__main__':unittest.main()
