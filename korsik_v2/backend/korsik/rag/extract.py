"""Text extraction for PDF / PPTX / DOCX / images with OCR fallback.

Used by both flows, which stay separate downstream:
  * corporate documents -> corpus (RAG);
  * personal employee documents -> private HR card fields.

Results are cached by file sha256 so repeated preparation does not redo OCR.
"""
from __future__ import annotations

import hashlib
import io
import json
import logging
import re
import zipfile
from dataclasses import dataclass, asdict
from pathlib import Path

log = logging.getLogger(__name__)

OCR_LANG = "rus+eng"
MIN_PAGE_TEXT = 40          # below this a PDF page is treated as a scan
MIN_OCR_IMAGE_PX = 250 * 250  # small pictures (icons/logos) are not OCR'ed


@dataclass
class Page:
    number: int            # 1-based page or slide number
    text: str
    method: str            # text | ocr | text+ocr | smartart
    notes: str = ""


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def _clean(text: str) -> str:
    text = text.replace("\xa0", " ").replace("​", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _ocr_image_bytes(data: bytes) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError:  # pragma: no cover
        return ""
    try:
        img = Image.open(io.BytesIO(data))
        if img.width * img.height < MIN_OCR_IMAGE_PX:
            return ""
        img = img.convert("L")
        if max(img.size) < 1600:  # upscale small screenshots for tesseract
            scale = 1600 / max(img.size)
            img = img.resize((int(img.width * scale), int(img.height * scale)))
        return _clean(pytesseract.image_to_string(img, lang=OCR_LANG))
    except Exception as exc:  # broken / unsupported picture
        log.warning("ocr image failed: %s", type(exc).__name__)
        return ""


def _meaningful_ocr(text: str) -> bool:
    words = re.findall(r"[А-Яа-яЁёA-Za-z]{3,}", text)
    return len(words) >= 4


# ---------------------------------------------------------------- PDF
def extract_pdf(path: Path) -> list[Page]:
    import fitz  # PyMuPDF
    pages: list[Page] = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc, 1):
            text = _clean(page.get_text("text"))
            if len(text) >= MIN_PAGE_TEXT:
                pages.append(Page(i, text, "text"))
                continue
            pix = page.get_pixmap(dpi=300)
            ocr = _ocr_image_bytes(pix.tobytes("png"))
            pages.append(Page(i, ocr or text, "ocr" if ocr else "text",
                              "" if ocr else "Страница не распознана"))
    return pages


# ---------------------------------------------------------------- PPTX
_A_T = re.compile(r"<a:t>([^<]*)</a:t>")


def _xml_text(xml: str) -> str:
    import html
    parts = [html.unescape(t) for t in _A_T.findall(xml)]
    return "\n".join(p for p in (s.strip() for s in parts) if p)


def extract_pptx(path: Path) -> list[Page]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    pages: list[Page] = []

    def walk(shapes, out_text, out_images):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(sh.shapes, out_text, out_images)
                continue
            if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
                out_text.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    out_text.append(" | ".join(c for c in cells if c))
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    out_images.append(sh.image.blob)
                except Exception:
                    pass

    for i, slide in enumerate(prs.slides, 1):
        texts: list[str] = []
        images: list[bytes] = []
        walk(slide.shapes, texts, images)
        method = "text"
        # SmartArt: text lives in related diagram data parts, not in shapes.
        for rel in slide.part.rels.values():
            if "diagramData" in rel.reltype:
                smart = _xml_text(rel.target_part.blob.decode("utf-8", "ignore"))
                if smart:
                    texts.append(smart)
                    method = "smartart"
        ocr_parts = []
        for blob in images:
            o = _ocr_image_bytes(blob)
            if o and _meaningful_ocr(o):
                ocr_parts.append(o)
        if ocr_parts:
            texts.append("\n".join(ocr_parts))
            method = method + "+ocr"
        pages.append(Page(i, _clean("\n".join(texts)), method))
    return pages


# ---------------------------------------------------------------- DOCX
def extract_docx(path: Path) -> list[Page]:
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            parts.append(" | ".join(c.text.strip() for c in r.cells if c.text.strip()))
    return [Page(1, _clean("\n".join(parts)), "text")]


def extract_image(path: Path) -> list[Page]:
    text = _ocr_image_bytes(path.read_bytes())
    return [Page(1, text, "ocr", "" if text else "Изображение не распознано")]


def extract_txt(path: Path) -> list[Page]:
    return [Page(1, _clean(path.read_text(encoding="utf-8", errors="ignore")), "text")]


EXTRACTORS = {
    ".pdf": extract_pdf, ".pptx": extract_pptx, ".docx": extract_docx,
    ".png": extract_image, ".jpg": extract_image, ".jpeg": extract_image,
    ".tif": extract_image, ".tiff": extract_image, ".txt": extract_txt,
}


def extract(path: Path, cache_dir: Path | None = None) -> list[Page]:
    ext = path.suffix.lower()
    if ext not in EXTRACTORS:
        raise ValueError(f"Неподдерживаемый формат: {ext}")
    cache_file = None
    if cache_dir:
        cache_dir.mkdir(parents=True, exist_ok=True)
        cache_file = cache_dir / f"{sha256_file(path)}.json"
        if cache_file.exists():
            return [Page(**p) for p in json.loads(cache_file.read_text(encoding="utf-8"))]
    pages = EXTRACTORS[ext](path)
    if cache_file:
        tmp = cache_file.with_suffix(".tmp")
        tmp.write_text(json.dumps([asdict(p) for p in pages], ensure_ascii=False), encoding="utf-8")
        tmp.replace(cache_file)
    return pages


if __name__ == "__main__":  # diagnostic: python -m korsik.rag.extract file.pdf
    import sys
    for p in extract(Path(sys.argv[1]), Path(sys.argv[2]) if len(sys.argv) > 2 else None):
        print(f"--- page {p.number} [{p.method}] {p.notes}\n{p.text}\n")
